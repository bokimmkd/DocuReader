from pathlib import Path
from zipfile import ZipFile
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
p=Path('app/src/androidTest/assets');p.mkdir(parents=True,exist_ok=True)
(p/'sample.csv').write_text('CSV_START;Amount;Кирилица\n"A;B";12;Табела\n'+'\n'.join(f'Line {i};{i};Текст' for i in range(80))+'\nCSV_END;99;Крај',encoding='utf-8')
(p/'sample.rtf').write_bytes(b'{\\rtf1\\ansi\\ansicpg1251 RTF_START\\par '+ 'Кирилица'.encode('cp1251')+b'\\par RTF_END}')
with ZipFile(p/'sample.epub','w') as z:
 z.writestr('mimetype','application/epub+zip')
 z.writestr('META-INF/container.xml','<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OPS/book.opf"/></rootfiles></container>')
 z.writestr('OPS/book.opf','<package xmlns="http://www.idpf.org/2007/opf"><manifest><item id="two" href="two.xhtml"/><item id="one" href="one.xhtml"/></manifest><spine><itemref idref="one"/><itemref idref="two"/></spine></package>')
 z.writestr('OPS/one.xhtml','<html><body><h1>EPUB_START Кирилица</h1><p>First chapter</p><img src="pic.png"/><script>BAD_SCRIPT</script><img src="https://example.com/pixel.png"/></body></html>')
 z.writestr('OPS/two.xhtml','<html><body><h1>Second chapter</h1><p>EPUB_END Крај</p></body></html>')
 z.writestr('OPS/pic.png',(p/'sample.png').read_bytes())
with ZipFile(p/'sample.odt','w') as z:
 z.writestr('content.xml','''<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" xmlns:xlink="http://www.w3.org/1999/xlink"><office:body><office:text><text:h>ODT_START Кирилица</text:h><text:p>Readable document</text:p><table:table><table:table-row><table:table-cell><text:p>Табела</text:p></table:table-cell></table:table-row></table:table><draw:frame><draw:image xlink:href="Pictures/pic.png"/></draw:frame><text:p>ODT_END Крај</text:p></office:text></office:body></office:document-content>''')
 z.writestr('Pictures/pic.png',(p/'sample.png').read_bytes())
r=Presentation()
for marker in ['PPTX_START Кирилица','PPTX_END Крај']:
 s=r.slides.add_slide(r.slide_layouts[5]);s.shapes.title.text=marker
 s.shapes.add_picture(str(p/'sample.png'),Inches(1),Inches(2))
 s.shapes.add_table(2,2,Inches(3),Inches(2),Inches(4),Inches(1)).table.cell(0,0).text='Табела'
r.save(p/'sample.pptx')
for ext in ['jpg','png','webp']:
 im=Image.new('RGB',(300,160),'#0869e8');im.save(p/f'picture.{ext}')
print('Extended-format fixtures created')

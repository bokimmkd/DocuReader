from pathlib import Path
from docx import Document
from openpyxl import Workbook
from PIL import Image
root=Path('app/src/androidTest/assets');root.mkdir(parents=True,exist_ok=True)
im=Image.new('RGB',(80,40),'#087cdb'); im.save(root/'sample.png')
doc=Document();doc.add_heading('WORD_START Кирилица',0)
doc.add_picture(str(root/'sample.png'))
t=doc.add_table(rows=3,cols=2);t.style='Table Grid'
for i,row in enumerate(t.rows):
    row.cells[0].text='Cell '+str(i);row.cells[1].text='Табела '+str(i)
for i in range(35):doc.add_paragraph('Word content '+str(i)+' with Unicode: Македонски текст')
doc.add_paragraph('Math ∈ ∉ ∪ ∩ ⊂ ≤ ≥ ℕ π Ω ∞');doc.add_page_break();doc.add_paragraph('WORD_END Крај')
doc.save(root/'word.docx')
wb=Workbook();ws=wb.active;ws.title='Invoices'
ws.append(['EXCEL_START','Qty','Price','Total','Status','Кирилица'])
for i in range(2,151):ws.append(['Item '+str(i),i,2.50,i*2.50,'Paid','Табела'])
ws['A150']='EXCEL_END';ws['C2'].number_format='$0.00'
for col in 'ABCDEF':ws.column_dimensions[col].width=16
wb.create_sheet('Do not print')['A1']='OTHER_SHEET_NOT_SELECTED'
wb.create_sheet('Too large')['A801']='TRUNCATED_END'
wb.save(root/'excel.xlsx')
(root/'text.txt').write_text('TXT_START <tag> Кирилица\n'+'\n'.join('Line '+str(i)+' Македонски текст' for i in range(150))+'\nMath ∈ ∉ ∪ ∩ ⊂ ≤ ≥ ℕ π Ω ∞\nTXT_END Крај',encoding='utf-8')
print('Print fixtures created')

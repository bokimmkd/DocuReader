#!/usr/bin/env python3
"""Create temporary QA Office/PDF fixtures using well-maintained test libraries.
Never ship these documents in the Android application."""
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

root = Path("verification/fixtures")
root.mkdir(parents=True, exist_ok=True)

doc = Document()
doc.add_heading("W Reader Word test", 0)
doc.add_paragraph("This paragraph tests rich DOCX rendering and scrolling.")
table = doc.add_table(rows=3, cols=2)
table.style = "Table Grid"
for row, a, b in [(0,"Product","Price"),(1,"Paper","12.50"),(2,"Toner","84.00")]:
    table.cell(row,0).text = a
    table.cell(row,1).text = b
for i in range(35):
    doc.add_paragraph("Document page content line %d, with bold: " % i).add_run("Office rendering").bold = True
doc.save(root / "word.docx")

wb = Workbook()
ws = wb.active
ws.title = "Budget"
ws["A1"] = "W Reader Excel test"
ws.merge_cells("A1:C1")
ws["A2"],ws["B2"],ws["C2"] = "Item","Qty","Total"
ws["A3"],ws["B3"],ws["C3"] = "Paper",2,12.50
ws["A4"],ws["B4"],ws["C4"] = "Toner",1,84.00
ws["C5"] = "=SUM(C3:C4)"
ws["C3"].number_format = "$#,##0.00"
ws["C4"].number_format = "$#,##0.00"
wb.create_sheet("Invoices")["A1"] = "Second sheet"
wb.save(root / "excel.xlsx")

pdf = canvas.Canvas(str(root / "pages.pdf"))
for i in range(1,7):
    pdf.setFont("Helvetica", 27)
    pdf.drawString(55,730,"W Reader page %d" % i)
    pdf.showPage()
pdf.save()
print("Created:", [(p.name,p.stat().st_size) for p in root.iterdir()])
